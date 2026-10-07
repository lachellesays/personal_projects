import io
import json
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from core import Show, parse_entries, transform, to_xlsx_bytes  # noqa: E402


def submission(sub_id, first, last, handler, dog_no, dog, sat=None, sun=None, paid='0.00',
               level='Novice', ss_level='Senior', height='20 regular', addr='1 Main St'):
    answers = {
        'name': {'first': first, 'last': last},
        'email': f'{first.lower()}@example.com',
        'address': {'addr_line1': addr, 'addr_line2': '', 'city': 'Richmond',
                    'state': 'Virginia', 'postal': '23220-1234'},
        'handlersUki': handler,
        'dogsName': dog,
        'dogsBreed': 'Mix',
        'dogsUki': dog_no,
        'jumpHeight': height,
        'internationalLevel': level,
        'speedstakesLevel': ss_level,
        'saturdayClasses': sat,
        'sundayClasses': sun,
        'myProducts': {'paymentArray': json.dumps({'total': paid})},
    }
    return {
        'id': sub_id, 'status': 'ACTIVE', 'created_at': '2026-06-01 10:00:00',
        'answers': {str(i): {'name': k, 'answer': v} for i, (k, v) in enumerate(answers.items())},
    }


SUBS = [
    submission('1', 'ann', 'smith', 'H100', 'D500', 'Rex',
               sat=['Agility 1', 'Masters Series'], sun=['Open Speedstakes 1 (all levels combined)'], paid='65.00'),
    submission('2', 'ann', 'smith', 'H100', 'D501', 'Bea', sat=['Jumping 1'], paid='15.00'),
    submission('3', 'Bob', 'Jones', '200', '600', 'Max', sun=['Snooker', 'Agility 2'], paid='10.00',
               level='Champion', height='16 select'),
]

SHOW = dict(name='Test', show_id=1234, form_id='f', saturday_date='2026-06-06', sunday_date='2026-06-07')


def run(**overrides):
    return transform(parse_entries(SUBS), Show(**{**SHOW, **overrides}))


def test_runs_levels_heights_and_masters_expansion():
    raw = run().raw_results
    ann_rex = raw[raw['Dog Name'] == 'Rex']
    assert sorted(ann_rex['Class']) == ['Agility', 'Masters Agility', 'Masters Jumping', 'Speedstakes']
    assert ann_rex.loc[ann_rex['Class'] == 'Speedstakes', 'Level'].item() == 'Senior'
    bob = raw[raw['Member Name'] == 'Bob Jones']
    assert set(bob['Level']) == {'Champ'}
    assert set(bob['Height']) == {'16 inch (s)'}
    assert raw['Member ID'].tolist().count(100) == 5
    assert len(raw) == 7


def test_balances_sum_payments_across_submissions():
    bal = run().balances.set_index('Member ID')['OWES']
    # Ann: 15 + 17.50*2 + 15 + 15 = 80, paid 65 + 15
    assert bal[100] == 0
    # Bob: 2 classes = 30, paid 10
    assert bal[200] == 20


def test_contact_formatting():
    contact = run().contact.set_index('Member ID')
    ann = contact.loc[100]
    assert (ann['First Name'], ann['Last Name']) == ('Ann', 'Smith')
    assert ann['County'] == 'VA'
    assert ann['Postcode'] == 23220
    assert ann['Address 1'] == '1 MAIN ST'
    assert len(contact) == 2


def test_scratched_handler_removed_from_all_tabs():
    r = run(scratched_handlers=['100'])
    assert 100 not in r.contact['Member ID'].tolist()
    assert 100 not in r.balances['Member ID'].tolist()
    assert 100 not in r.raw_results['Member ID'].tolist()
    assert len(r.raw_results) == 2


def test_scratched_dog_keeps_handlers_other_dog():
    r = run(scratched_dogs=['500'])
    assert set(r.raw_results.loc[r.raw_results['Member ID'] == 100, 'Dog Name']) == {'Bea'}
    # Bea's one class (15) minus everything Ann paid (80)
    assert r.balances.set_index('Member ID')['OWES'][100] == -65
    assert any('negative balance' in w for w in r.warnings)


def test_exclude_addresses_picks_other_variant():
    subs = SUBS + [submission('4', 'Bob', 'Jones', '200', '601', 'Zip', sat=['Agility 1'], addr='1 Main Street')]
    subs[2] = submission('3', 'Bob', 'Jones', '200', '600', 'Max', sun=['Snooker'], addr='1 MAIN ST')
    show = Show(**SHOW, exclude_addresses=['1 main st'])
    r = transform(parse_entries(subs), show)
    assert r.contact.set_index('Member ID').loc[200, 'Address 1'] == '1 MAIN STREET'


def test_missing_field_raises_clear_error():
    broken = [submission('1', 'a', 'b', '1', '2', 'x')]
    for a in broken[0]['answers'].values():
        if a['name'] == 'handlersUki':
            a['name'] = 'handlerNumber'
    with pytest.raises(ValueError, match='handlersUki'):
        parse_entries(broken)
    assert parse_entries(broken, {'handler_number': 'handlerNumber'})[0]['handler_number'] == '1'


def test_xlsx_layout_matches_template():
    wb = openpyxl.load_workbook(io.BytesIO(to_xlsx_bytes(run())))
    assert wb.sheetnames == ['Contact', 'Balances', 'Raw Results']
    ws = wb['Raw Results']
    header = [c.value for c in ws[1]]
    row = dict(zip(header, ws[2]))
    assert row['Date'].number_format == 'yyyy-mm-dd'
    assert row['Date'].is_date
    assert isinstance(row['Member ID'].value, int)
    assert row['Class'].number_format == '@'
    assert row['SCT'].value is None


def test_corrections_apply_to_every_tab():
    from core import Show as S
    show = S(**SHOW, handler_corrections={'100': {'last_name': 'Smyth', 'state': 'ALABAMA'}},
             dog_corrections={'600': {'dog_name': 'Maxine', 'jump_height': '12 regular'}})
    r = transform(parse_entries(SUBS), show)
    c = r.contact.set_index('Member ID')
    assert c.loc[100, 'Last Name'] == 'Smyth'
    assert c.loc[100, 'County'] == 'AL'
    assert r.balances.set_index('Member ID').loc[100, 'Last Name'] == 'Smyth'
    assert set(r.raw_results.loc[r.raw_results['Member ID'] == 100, 'Member Name']) == {'ann Smyth'}
    bob = r.raw_results[r.raw_results['Member ID'] == 200]
    assert set(bob['Dog Name']) == {'Maxine'} and set(bob['Height']) == {'12 inch'}


def test_diff_corrections_and_yaml_round_trip():
    from core import handler_table, diff_corrections, HANDLER_FIELDS, Show as S
    entries = parse_entries(SUBS)
    orig = handler_table(entries)
    edited = orig.copy()
    edited.loc[edited['handler_number'] == '200', 'city'] = 'Norfolk'
    fixes = diff_corrections(orig, edited, 'handler_number', HANDLER_FIELDS)
    assert fixes == {'200': {'city': 'Norfolk'}}
    show = S(**SHOW, handler_corrections=fixes)
    assert S.from_dict(show.to_dict()).handler_corrections == fixes
