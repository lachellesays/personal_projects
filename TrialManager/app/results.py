"""
results.py

Scoring a trial: course details (yardage), score entry, Calculate (SCT, faults, Q, places, level points),
and the outputs — results JSON for the check-in app's "Publish results" and the UKI results spreadsheet.
The rules themselves live in scoring.py.
"""

import io
from datetime import date
from pathlib import Path

from sqlalchemy import delete, insert, select, update

from app import db, scoring, trial

FAULT_CHOICES = [0, 5, 10, 15, 20, 25, 30, 35, 40, 45, 50]
REFUSAL_CHOICES = [0, 5, 10, 15, 20]
GAMBLE_CHOICES = [0, 10, 15, 20]
EMPTY_SCORE = {'status': '', 'time': None, 'course_faults': None, 'refusals': None, 'opening_points': None, 'opening_path': '',
               'gamble_points': None, 'points': None, 'result': None}
COURSE_FIELDS = ('yards_bn', 'yards_senior', 'yards_champ', 'sct_extra', 'opening_time', 'closing_time', 'course_time')


class ResultsError(Exception):
    pass


def course_of(setting: dict) -> dict:
    return {k: setting.get(k) for k in COURSE_FIELDS + ('rate',)}


def _int_or_none(v, lo=0, hi=999):
    v = str(v if v is not None else '').strip()
    if v == '':
        return None
    if not v.isdigit() or not lo <= int(v) <= hi:
        raise ResultsError(f'"{v}" isn’t a valid number here.')
    return int(v)


def save_course(slug: str, day: str, class_name: str, form: dict, user: dict):
    setting = trial.get_class(slug, day, class_name)
    values = {k: _int_or_none(form.get(k)) for k in COURSE_FIELDS}
    values['sct_extra'] = values['sct_extra'] or 0
    if values['sct_extra'] > 3:
        raise ResultsError('The judge can add at most 3 seconds.')
    values['rate'] = 'slow' if form.get('rate') == 'slow' else 'fast'
    with db.engine().begin() as c:
        c.execute(update(db.class_settings).where(db.class_settings.c.id == setting['id']).values(calculated_at=None, **values))
    db.log(user, 'course', slug, target=f'{day}|{class_name}', **values)


# ── Score entry ─────────────────────────────────────────────────────────────

def _scores(slug: str, run_ids: list) -> dict:
    if not run_ids:
        return {}
    with db.engine().connect() as c:
        return {r['run_id']: dict(r) for r in c.execute(select(db.scores).where(db.scores.c.run_id.in_(run_ids))).mappings()}


def class_runs(slug: str, day: str, class_name: str) -> list:
    """Active runs in running order, each with its score row (if any)."""
    runs_ = [r for r in trial.ordered_for_export(slug, day) if r['class_name'] == class_name]
    sc = _scores(slug, [r['id'] for r in runs_])
    return [{**r, 'score': sc.get(r['id'])} for r in runs_]


def _num(v, field):
    v = str(v if v is not None else '').strip()
    if v == '':
        return None
    try:
        n = float(v)
    except ValueError:
        raise ResultsError(f'{field} should be a number, like 32.15.')
    if n < 0 or n > 999:
        raise ResultsError(f'{field} looks wrong: {v}.')
    return round(n, 2)


def save_score(slug: str, run_id: int, form: dict, user: dict):
    with db.engine().connect() as c:
        run = c.execute(select(db.runs).where(db.runs.c.id == run_id, db.runs.c.show_slug == slug)).mappings().first()
    if not run:
        raise ResultsError('That run doesn’t exist.')
    status = (form.get('status') or '').upper()
    if status not in ('',) + scoring.STATUSES:
        raise ResultsError('Status must be blank, E, ABS or NFC.')
    values = {'status': status, 'time': _num(form.get('time'), 'Time'),
              'course_faults': _num(form.get('course_faults'), 'Faults'), 'refusals': _num(form.get('refusals'), 'Refusals'),
              'gamble_points': _int_or_none(form.get('gamble_points')), 'points': _int_or_none(form.get('points'))}
    try:
        values['opening_path'], values['opening_points'] = scoring.opening_points_from(form.get('opening_points'))
    except ValueError as e:
        raise ResultsError(str(e))
    with db.engine().begin() as c:
        c.execute(delete(db.scores).where(db.scores.c.run_id == run_id))
        c.execute(insert(db.scores).values(run_id=run_id, show_slug=slug, result=None, updated_at=db.now(),
                                           updated_by=user.get('name', ''), **values))
        c.execute(update(db.class_settings).where(db.class_settings.c.show_slug == slug, db.class_settings.c.day == run['day'],
                                                  db.class_settings.c.class_name == run['class_name']).values(calculated_at=None))
    db.log(user, 'score', slug, target=f'r:{run_id}', **values)


def has_score(score: dict) -> bool:
    return bool(score) and (bool(score.get('status')) or score.get('time') is not None or score.get('points') is not None
                            or score.get('opening_points') is not None)


def _yards(run: dict, course: dict, class_name: str):
    if scoring.is_open_class(class_name, course) or run['level'] not in ('Beginner', 'Novice', 'Nursery'):
        return (course.get('yards_senior') if run['level'] == 'Senior' else None) or course.get('yards_champ')
    return course.get('yards_bn')


def _masters_forfeits(slug: str, day: str, class_name: str) -> set:
    """Dogs absent/NFC in the other Masters round of the day: they forfeit Q's and awards in this one."""
    if class_name not in MASTERS_CLASSES:
        return set()
    other = MASTERS_CLASSES[1 - MASTERS_CLASSES.index(class_name)]
    if other not in [s['class_name'] for s in trial.class_list(slug, day)]:
        return set()
    return {_dog_key(r) for r in class_runs(slug, day, other) if (r['score'] or {}).get('status') in ('ABS', 'NFC')}


def calculate(slug: str, day: str, class_name: str, user: dict, _partner: bool = True) -> dict:
    """Work out SCT, faults, Q, places and level points for one class and save them."""
    setting = trial.get_class(slug, day, class_name)
    forfeits = _masters_forfeits(slug, day, class_name)
    course = course_of(setting)
    runs_ = class_runs(slug, day, class_name)
    missing_sct = 0
    for r in runs_:
        sct_ = scoring.sct(class_name, r['level'], r['height'], r['class_type'], course)
        missing_sct += sct_ is None
        r['result'] = scoring.score_run(class_name, r['score'] or {}, sct_, course, _yards(r, course, class_name))
        r['result']['scored'] = has_score(r['score'])
        if not r['result']['scored']:
            r['result'].update(q=False, total_faults=None, time_faults=0.0, yps=None, points=None)
        if _dog_key(r) in forfeits:
            r['result'].update(q=False, forfeit=True)
    if missing_sct:
        raise ResultsError('Fill in the course details (yardage or times) before calculating.')
    scoring.place_class(class_name, runs_, course)
    with db.engine().begin() as c:
        for r in runs_:
            result = {**r['result'], 'place': r['place'], 'level_points': r['level_points']}
            if r['score']:
                c.execute(update(db.scores).where(db.scores.c.run_id == r['id']).values(result=result))
            else:
                c.execute(insert(db.scores).values(run_id=r['id'], show_slug=slug, status='', result=result,
                                                   updated_at=db.now(), updated_by=user.get('name', '')))
        c.execute(update(db.class_settings).where(db.class_settings.c.id == setting['id']).values(calculated_at=db.now()))
    unscored = sum(1 for r in runs_ if not r['result']['scored'])
    db.log(user, 'calculate', slug, target=f'{day}|{class_name}', unscored=unscored)
    if _partner and class_name in MASTERS_CLASSES:   # an absence here can change the other round's awards
        other = MASTERS_CLASSES[1 - MASTERS_CLASSES.index(class_name)]
        if any(s['class_name'] == other and s['calculated_at'] for s in trial.class_list(slug, day)):
            calculate(slug, day, other, user, _partner=False)
    return {'runs': len(runs_), 'unscored': unscored, 'qs': sum(1 for r in runs_ if r['result']['q'])}


# ── Views ───────────────────────────────────────────────────────────────────

def overview(slug: str, day: str) -> list:
    out = []
    for s in trial.class_list(slug, day):
        runs_ = class_runs(slug, day, s['class_name'])
        out.append({**s, 'runs': len(runs_), 'scored': sum(1 for r in runs_ if has_score(r['score'])),
                    'calculated': bool(s['calculated_at'])})
    return out


def entry_view(slug: str, day: str, class_name: str) -> dict:
    setting = trial.get_class(slug, day, class_name)
    course = course_of(setting)
    rows = []
    for r in class_runs(slug, day, class_name):
        r['score'] = {**EMPTY_SCORE, **(r['score'] or {})} if r['score'] else None
        rows.append({**r, 'sct': scoring.sct(class_name, r['level'], r['height'], r['class_type'], course),
                     'result': (r['score'] or {}).get('result') if setting['calculated_at'] else None})
    return {'setting': setting, 'course': course, 'kind': scoring.class_kind(class_name), 'rows': rows,
            'open': scoring.is_open_class(class_name, course), 'calculated': bool(setting['calculated_at']),
            'scored': sum(1 for r in rows if has_score(r['score']))}


def results_view(slug: str, day: str, class_name: str) -> list:
    """Calculated results grouped the way they're placed (level, height, Regular/Select)."""
    ev = entry_view(slug, day, class_name)
    groups = {}
    for r in ev['rows']:
        res = r['result'] or {}
        key = scoring.placement_group(class_name, r['level'], r['height'], r['class_type'], ev['course'])
        groups.setdefault(key, []).append({**r, 'res': res})
    out = []
    for key, rows in groups.items():
        rows.sort(key=lambda r: (r['res'].get('place') or 999, r['res'].get('status') or '', -(r['res'].get('points') or 0),
                                 r['res'].get('time') or 999))
        label = (f'{key[0]} · ' if len(key) == 3 else '') + f'{key[-2]}" {key[-1]}'
        out.append({'label': label, 'key': key, 'rows': rows})
    out.sort(key=lambda g: (str(g['key'][0]) if len(g['key']) == 3 else '', g['key'][-2], g['key'][-1]))
    return out


# ── Outputs ─────────────────────────────────────────────────────────────────

def _fmt(v):
    return '' if v is None else (f'{v:.2f}' if isinstance(v, float) else str(v))


def results_json(slug: str, day: str, class_name: str, combined: str = None) -> list:
    """One class's results in UKI Manager's JSON layout (what the check-in app's Publish results reads).
    `combined` limits it to one check-in class, e.g. 'Senior/Champ Agility'."""
    ev = entry_view(slug, day, class_name)
    if not ev['calculated']:
        raise ResultsError('Calculate this class first.')
    out = []
    for r in ev['rows']:
        if combined and r['combined'] != combined:
            continue
        res, s = r['result'] or {}, r['score'] or {}
        if not res.get('scored'):
            continue   # no score yet: left out rather than shown as a clear round
        status = res.get('status') or ''
        item = {'class_type': r['class_type'], 'height': str(r['height']),
                'uki_number': f"{r['first_name']} {r['last_name']}".strip(), 'uki_dog_number': r['dog_name'],
                'sct': _fmt(res.get('sct')), 'time': f"{res.get('time') or 0:.2f}",
                'place': str(res['place']) if res.get('place') else ' ',
                'level_points': str(res['level_points']) if res.get('level_points') else ' ',
                'qualify': 'Y' if res.get('q') else ' '}
        if ev['kind'] == 'gamblers':
            item.update(faults=status or '0', games_points=str(s.get('opening_points') or 0),
                        gamble_points=str(s.get('gamble_points') or 0), total_points=res.get('points') or 0)
        elif ev['kind'] == 'snooker':
            item.update(faults=status or '0', total_points=res.get('points') or 0)
        else:
            item.update(faults=status or _fmt(res.get('total_faults') if res.get('total_faults') is not None else 0),
                        faults_standard=status or _fmt(s.get('course_faults') or 0), faults_refusal=_fmt(s.get('refusals') or 0),
                        timefaults=_fmt(res.get('time_faults')), yps=_fmt(res.get('yps')))
        out.append(item)
    return out


RAW_RESULTS_COLUMNS = ['Show ID', 'Level', 'Class', 'Height', 'Date', 'Member ID', 'Member Name', 'Dog ID', 'Dog Name',
                       'Breed', 'SCT', 'Faults', 'Games Points', 'Time', 'Place', 'Level Points']


def results_workbook(slug: str, show) -> bytes:
    """The Raw Results sheet in UKI Manager's results spreadsheet layout, for sending to UKI."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Raw Results'
    ws.append(RAW_RESULTS_COLUMNS)
    for day in (str(show.saturday_date), str(show.sunday_date)):
        for s in trial.class_list(slug, day):
            if not s['calculated_at']:
                continue
            for r in class_runs(slug, day, s['class_name']):
                res, sc_ = (r['score'] or {}).get('result') or {}, r['score'] or {}
                status = res.get('status') or ''
                faults = status or (int(sc_.get('course_faults') or 0) + int(sc_.get('refusals') or 0) if res.get('scored') else '')
                ws.append([show.show_id, r['level'], s['class_name'],
                           f"{r['height']} inch" + (' (s)' if r['class_type'] == 'Select' else ''),
                           date.fromisoformat(day).strftime('%-m/%-d/%Y'), _intish(r['handler_number']),
                           f"{r['first_name']} {r['last_name']}".strip(), _intish(r['dog_number']), r['dog_name'], r['breed'],
                           res.get('sct'), faults, res.get('points') or 0, round(res.get('time') or 0, 2),
                           res.get('place') or 0, res.get('level_points') or 0])
    _masters_sheets(wb, slug, show)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _masters_sheets(wb, slug: str, show):
    """'Masters Scoring' and 'Masters Byes' sheets, like UKI Manager's workbook, when Masters ran."""
    for day in (str(show.saturday_date), str(show.sunday_date)):
        mv = masters_view(slug, day)
        if not mv['available']:
            continue
        ws = wb.create_sheet('Masters Scoring')
        ws.append(['Height', 'HandlerID', 'HandlerName', 'DogID', 'DogName', 'Jumping SCT', 'Jumping Faults', 'Jumping Time',
                   'Jumping Total', 'Agility SCT', 'Agility Faults', 'Agility Time', 'Agility Total', 'Overall Total',
                   'Place', 'Prize $', 'Final Qualified'])
        byes = wb.create_sheet('Masters Byes')
        byes.append(['Jump Height', 'Handler #', 'Handler Name', 'Dog #', 'Dog Name'])
        for g in mv['groups']:
            height = f"{g['height']} inch" + ('(s)' if g['class_type'] == 'Select' else '')
            for d in g['rows']:
                ws.append([height, _intish(d['handler_number']), d['handler'], _intish(d['dog_number']), d['dog_name']]
                          + _round_cells(d.get('jumping'), d.get('jumping_sct')) + [d['jumping_score']]
                          + _round_cells(d.get('agility'), d.get('agility_sct')) + [d['agility_score']]
                          + [d['overall'], d['place'], d['prize'], 'Y' if d['bye'] else ''])
                if d['bye']:
                    byes.append([height, _intish(d['handler_number']), d['handler'], _intish(d['dog_number']), d['dog_name']])
        return


def _intish(v):
    return int(v) if str(v or '').isdigit() else v


# ── Masters Series (both rounds combined) ───────────────────────────────────

MASTERS_CLASSES = ('Masters Agility', 'Masters Jumping')


def _dog_key(r: dict) -> str:
    return r['dog_number'] or f"{r['handler_number']}:{r['dog_name'].lower()}"


def masters_view(slug: str, day: str) -> dict:
    """Overall Masters Series results for one day, or why they can't be worked out yet."""
    names = [s['class_name'] for s in trial.class_list(slug, day)]
    if not all(c in names for c in MASTERS_CLASSES):
        return {'available': False, 'reason': 'This day doesn’t have both Masters Agility and Masters Jumping.'}
    views = {c: entry_view(slug, day, c) for c in MASTERS_CLASSES}
    waiting = [c for c, v in views.items() if not v['calculated']]
    if waiting:
        return {'available': False, 'reason': 'Calculate ' + ' and '.join(waiting) + ' first.'}
    dogs = {}
    for cls, key in (('Masters Agility', 'agility'), ('Masters Jumping', 'jumping')):
        for r in views[cls]['rows']:
            d = dogs.setdefault(_dog_key(r), {'id': _dog_key(r), 'dog_name': r['dog_name'], 'dog_number': r['dog_number'],
                                              'handler': f"{r['first_name']} {r['last_name']}".strip(),
                                              'handler_number': r['handler_number'], 'height': r['height'],
                                              'class_type': r['class_type']})
            d[key] = r['result']
            d[key + '_sct'] = r['sct']
    rows = scoring.masters_series(list(dogs.values()))
    groups = {}
    for d in rows:
        groups.setdefault((d['height'], d['class_type']), []).append(d)
    out = []
    for (h, t), members in sorted(groups.items()):
        members.sort(key=lambda d: (d['place'] or 999, d['overall'] if d['overall'] is not None else 9999))
        out.append({'label': f'{h}" {t}', 'height': h, 'class_type': t, 'entered': len(members), 'rows': members,
                    'fund': scoring.masters_prizes(len(members))})
    return {'available': True, 'groups': out}


def _round_cells(res: dict, sct_):
    """SCT, faults and time for a round, the way UKI's Masters sheet shows them."""
    res = res or {}
    status = res.get('status') or ''
    if status == 'E':
        return [sct_, 100, 60]
    if status in ('ABS', 'NFC') or not res.get('scored'):
        return [sct_, 1000, 60]
    return [sct_, res.get('total_faults') or 0, round(res.get('time') or 0, 2)]


# ── Masters Final (bye) certificates ────────────────────────────────────────

CERTIFICATE_BACKGROUND = Path(__file__).parent / 'static' / 'certificates' / 'masters_final_2027.png'


def masters_certificates_pdf(slug: str, day: str, show_name: str) -> bytes:
    """UKI's Masters Final certificate for every qualifier that day, filled in like UKI Manager's
    (handler, dog, height, trial name & date). The Show manager line is left to sign."""
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    mv = masters_view(slug, day)
    if not mv['available']:
        raise ResultsError(mv['reason'])
    qualifiers = [d for g in mv['groups'] for d in g['rows'] if d['bye']]
    if not qualifiers:
        raise ResultsError('No dogs qualified for the Final that day.')
    pw, ph = landscape(letter)                     # 792 × 612, the template's shape
    sx, sy = pw / 1680, ph / 1300                  # template pixels → points
    def at(px, py):
        return px * sx, ph - py * sy
    background = ImageReader(str(CERTIFICATE_BACKGROUND))
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pagesize=(pw, ph))
    pdf.setTitle('Masters Final certificates')
    trial = f'{show_name}   on   {day}'
    for d in qualifiers:
        pdf.drawImage(background, 0, 0, width=pw, height=ph)
        height = f"{d['height']} inch" + ('(s)' if d['class_type'] == 'Select' else '')
        pdf.setFont('Times-Bold', 15)
        pdf.drawString(*at(290, 586), _fit(pdf, d['handler'], 'Times-Bold', 15, 530 * sx))
        pdf.drawString(*at(925, 586), _fit(pdf, d['dog_name'], 'Times-Bold', 15, 320 * sx))
        pdf.setFont('Times-Roman', 14)
        pdf.drawString(*at(1352, 586), height)
        size = 14
        while size > 9 and pdf.stringWidth(trial, 'Times-Roman', size) > 620 * sx:
            size -= 0.5
        pdf.setFont('Times-Roman', size)
        pdf.drawCentredString(*at(1146, 833), trial)
        pdf.showPage()
    pdf.save()
    return buf.getvalue()


def _fit(pdf, text: str, font: str, size: float, width: float) -> str:
    """Shrink the font so a long name stays on its line."""
    while size > 9 and pdf.stringWidth(text, font, size) > width:
        size -= 0.5
    pdf.setFont(font, size)
    return text
